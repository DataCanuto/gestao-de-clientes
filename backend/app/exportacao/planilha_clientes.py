"""Relatório .xlsx dos clientes: a "planilha-relatório" que o escritório usa.

Substitui a saída do notebook (tabela_final_clientes.xlsx e
clientes_arquivos_faltantes.xlsx, células 36 e 37), agora gerada a partir dos
dados CONFERIDOS no banco, e não de uma execução do notebook.

Abas:
- Clientes:   uma linha por cliente, com as colunas pedidas pela cliente
              (Nome, Nº do Processo, Corretor, Situação, Data de Emissão da NF)
              + quantidade de cada documento, o que falta e as pendências;
- Faltantes:  só quem ainda não tem os três documentos (o que buscar);
- Documentos: um arquivo por linha (para auditoria), inclusive os sem cliente;
- Resumo:     os números do painel.

Por que openpyxl, e não pandas? Só é preciso ESCREVER uma planilha com
formatação simples; o pandas traria uma dependência grande (e o numpy) para
isso. O openpyxl já era usado pelo script de paridade.
"""

from collections.abc import Iterable
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.entities import (
    DadosNotaFiscal,
    DadosPlanilhaComissao,
    DadosPrestacaoContas,
    DocumentoArmazenado,
    FichaCliente,
    ResultadoExtracao,
)
from app.domain.enums import TipoDocumento
from app.domain.relatorio import clientes_com_documentos_faltantes, gerar_resumo
from app.domain.rotulos import ROTULO_TIPO, rotulo_campos, rotulo_situacao

SEPARADOR = " | "  # vários valores numa célula (cliente com duas NFs, dois processos...)
FORMATO_DATA = "DD/MM/YYYY"
LARGURA_MAXIMA = 60

_CABECALHO_FONTE = Font(bold=True, color="FFFFFF")
_CABECALHO_FUNDO = PatternFill("solid", fgColor="1F4E78")

COLUNAS_CLIENTES = [
    "Nome do Cliente",
    "Espólio",
    "Nº do Processo",
    "Nome do Corretor",
    "Situação",
    "Data Emissão NF",
    "NFs",
    "Planilhas",
    "Prestações",
    "Documentos faltantes",
    "Completo",
    "Documentos a verificar",
]
COLUNAS_DOCUMENTOS = [
    "Cliente",
    "Tipo",
    "Arquivo",
    "Data Emissão NF",
    "Situação",
    "Corretor",
    "Nº do Processo",
    "Nome na Prestação",
    "Campos a verificar",
    "Corrigidos à mão",
]


def gerar_planilha_clientes(fichas: list[FichaCliente], documentos_sem_cliente: list[DocumentoArmazenado]) -> bytes:
    livro = Workbook()
    aba_clientes = livro.active
    aba_clientes.title = "Clientes"
    _preencher(aba_clientes, COLUNAS_CLIENTES, (_linha_cliente(f) for f in fichas))
    _preencher(
        livro.create_sheet("Faltantes"), COLUNAS_CLIENTES, (_linha_cliente(f) for f in clientes_com_documentos_faltantes(fichas))
    )
    linhas_documentos = [_linha_documento(f.nome, d) for f in fichas for d in f.documentos]
    linhas_documentos += [_linha_documento("(sem cliente)", d) for d in documentos_sem_cliente]
    _preencher(livro.create_sheet("Documentos"), COLUNAS_DOCUMENTOS, linhas_documentos)
    _preencher_resumo(livro.create_sheet("Resumo"), fichas, len(documentos_sem_cliente))

    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


# --------------------------------------------------------------------------- linhas


def _linha_cliente(ficha: FichaCliente) -> list:
    quantidade = {tipo: len(ficha.documentos_do_tipo(tipo)) for tipo in TipoDocumento}
    datas = ficha.datas_nf
    return [
        ficha.nome,
        "Sim" if ficha.espolio else "",
        SEPARADOR.join(ficha.numeros_processo),
        SEPARADOR.join(ficha.corretores),
        SEPARADOR.join(rotulo_situacao(s) for s in ficha.situacoes),
        # Uma data só vira DATA no Excel (ordenável, filtrável); várias viram texto.
        datas[0] if len(datas) == 1 else SEPARADOR.join(d.strftime("%d/%m/%Y") for d in datas),
        quantidade[TipoDocumento.NOTA_FISCAL],
        quantidade[TipoDocumento.PLANILHA_COMISSAO],
        quantidade[TipoDocumento.PRESTACAO_CONTAS],
        ", ".join(ROTULO_TIPO[tipo] for tipo in ficha.documentos_faltantes),
        "Sim" if ficha.completo else "Não",
        len(ficha.documentos_a_verificar),
    ]


def _linha_documento(cliente: str, documento: ResultadoExtracao) -> list:
    dados = documento.dados
    return [
        cliente,
        ROTULO_TIPO.get(documento.tipo, ""),
        documento.nome_arquivo,
        dados.data_emissao if isinstance(dados, DadosNotaFiscal) else None,
        rotulo_situacao(getattr(dados, "situacao", None)),
        dados.corretor if isinstance(dados, DadosPlanilhaComissao) else "",
        dados.numero_processo if isinstance(dados, DadosPrestacaoContas) else "",
        dados.nome_encontrado if isinstance(dados, DadosPrestacaoContas) else "",
        rotulo_campos(documento.campos_a_verificar),
        rotulo_campos(getattr(documento, "campos_editados", [])),
    ]


def _preencher_resumo(aba: Worksheet, fichas: list[FichaCliente], sem_cliente: int) -> None:
    resumo = gerar_resumo(fichas)
    linhas = [
        ["Total de clientes", resumo.total_clientes],
        ["Clientes completos (NF + Planilha + Prestação)", resumo.clientes_completos],
        ["Clientes com documentos faltantes", resumo.clientes_com_documentos_faltantes],
        *[[f"Clientes sem {ROTULO_TIPO[t]}", resumo.clientes_sem_documento[t]] for t in TipoDocumento],
        *[[f"Total de {ROTULO_TIPO[t]}", resumo.total_documentos[t]] for t in TipoDocumento],
        ["Documentos a verificar", resumo.documentos_a_verificar],
        ["Documentos sem cliente", sem_cliente],
        ["Gerado em", date.today()],
    ]
    _preencher(aba, ["Indicador", "Valor"], linhas)


# --------------------------------------------------------------------------- formatação


def _preencher(aba: Worksheet, colunas: list[str], linhas: Iterable[list]) -> None:
    aba.append(colunas)
    for celula in aba[1]:
        celula.font, celula.fill = _CABECALHO_FONTE, _CABECALHO_FUNDO
        celula.alignment = Alignment(vertical="center")
    for linha in linhas:
        aba.append(linha)
        for celula in aba[aba.max_row]:
            if isinstance(celula.value, date):
                celula.number_format = FORMATO_DATA

    aba.freeze_panes = "A2"  # cabeçalho sempre visível ao rolar
    aba.auto_filter.ref = aba.dimensions  # setas de filtro em cada coluna
    for indice, coluna in enumerate(aba.iter_cols(values_only=True), start=1):
        maior = max((len(_texto(v)) for v in coluna), default=10)
        aba.column_dimensions[get_column_letter(indice)].width = min(max(maior + 2, 10), LARGURA_MAXIMA)


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, date):
        return "00/00/0000"
    return str(valor)
