"""Arquivos gerados para pessoas: planilha .xlsx, PDF mesclado e as pastas organizadas."""

import json
from datetime import date
from io import BytesIO
from pathlib import Path

import pymupdf
from openpyxl import load_workbook

from app.core.armazenamento import ArmazenamentoDeArquivos
from app.core.libpq import localizar_pasta_da_libpq
from app.domain.entities import FichaCliente
from app.domain.enums import Situacao
from app.exportacao.pdf import mesclar_pdfs
from app.exportacao.planilha_clientes import gerar_planilha_clientes
from app.services.organizacao_service import NOME_MANIFESTO, OrganizacaoService
from tests.conftest import criar_pdf_com_texto
from tests.domain.test_organizacao import documento, ficha, nf, planilha, prestacao


class TestPlanilhaClientes:
    def test_colunas_pedidas_e_abas(self):
        ana = ficha(nf(), planilha(situacao=Situacao.ACORDO, corretor="BETO"), prestacao())
        bruno = FichaCliente(nome="BRUNO LIMA", chave="BRUNO LIMA", documentos=[nf(nome="BRUNO LIMA")], id=2)

        livro = load_workbook(BytesIO(gerar_planilha_clientes([ana, bruno], [])))

        assert livro.sheetnames == ["Clientes", "Faltantes", "Documentos", "Resumo"]
        clientes = list(livro["Clientes"].iter_rows(values_only=True))
        assert clientes[0][:6] == ("Nome do Cliente", "Espólio", "Nº do Processo", "Nome do Corretor", "Situação", "Data Emissão NF")
        linha_ana = clientes[1]
        assert linha_ana[:5] == ("ANA SOUZA", None, "0001", "BETO", "Acordo")
        assert linha_ana[5].date() == date(2023, 3, 8)  # data de verdade no Excel, e não texto
        assert linha_ana[10] == "Sim"  # completo
        faltantes = list(livro["Faltantes"].iter_rows(values_only=True))
        assert [linha[0] for linha in faltantes[1:]] == ["BRUNO LIMA"]
        assert faltantes[1][9] == "Planilha Comissão, Prestação de Contas"

    def test_documentos_sem_cliente_aparecem_na_aba_documentos(self):
        sem_cliente = nf(cliente_id=None, nome=None, nome_arquivo="NF 999.pdf")
        livro = load_workbook(BytesIO(gerar_planilha_clientes([], [sem_cliente])))
        linhas = list(livro["Documentos"].iter_rows(values_only=True))
        assert linhas[1][:3] == ("(sem cliente)", "Nota Fiscal", "NF 999.pdf")


def test_mesclar_pdfs_mantem_a_ordem_das_paginas(tmp_path):
    caminhos = []
    for texto in ("prestacao", "planilha", "nf"):
        caminho = tmp_path / f"{texto}.pdf"
        caminho.write_bytes(criar_pdf_com_texto([texto]))
        caminhos.append(caminho)

    with pymupdf.open(stream=mesclar_pdfs(caminhos), filetype="pdf") as mesclado:
        assert [pagina.get_text().strip() for pagina in mesclado] == ["prestacao", "planilha", "nf"]


class _Repositorio:
    """Dublê dos repositórios: o OrganizacaoService só precisa de listar()."""

    def __init__(self, fichas=None, sem_cliente=None):
        self._fichas, self._sem_cliente = fichas or [], sem_cliente or []

    def listar(self, *args, sem_cliente=False, **kwargs):
        return self._sem_cliente if sem_cliente else self._fichas


class TestOrganizacaoService:
    def _servico(self, tmp_path, fichas, sem_cliente=None) -> OrganizacaoService:
        armazenamento = ArmazenamentoDeArquivos(tmp_path / "storage")
        for ficha_ in fichas:
            for doc in ficha_.documentos:
                doc.hash_sha256 = armazenamento.salvar(criar_pdf_com_texto([doc.nome_arquivo, str(doc.id)]))
        repositorio = _Repositorio(fichas, sem_cliente)
        return OrganizacaoService(repositorio, repositorio, armazenamento, tmp_path / "saida")

    def test_gera_pastas_pdf_mesclado_relatorio_e_manifesto(self, tmp_path):
        servico = self._servico(tmp_path, [ficha(nf(), planilha(), prestacao())])

        resultado = servico.organizar()

        saida = tmp_path / "saida"
        assert (resultado.documentos_copiados, resultado.pdfs_mesclados, resultado.pendencias) == (3, 1, [])
        assert (saida / "Documentos Verificados" / "N Fiscais" / "NF Ana Souza.pdf").exists()
        with pymupdf.open(saida / "Documentos Mesclados" / "Ana Souza.pdf") as mesclado:
            assert mesclado.page_count == 3
        assert (saida / "Relatório de Clientes.xlsx").exists()
        assert len(json.loads((saida / NOME_MANIFESTO).read_text(encoding="utf-8"))["arquivos"]) == 5

    def test_remove_apenas_o_que_o_sistema_gerou_e_ficou_obsoleto(self, tmp_path):
        ana = ficha(nf(), planilha(), prestacao())
        servico = self._servico(tmp_path, [ana])
        servico.organizar()
        saida = tmp_path / "saida"
        (saida / "Documentos Mesclados" / "anotacoes pessoais.txt").write_text("não apagar")

        ana.nome = "ANA SOUZA SANTOS"  # cliente renomeado no sistema
        resultado = servico.organizar()

        assert resultado.arquivos_removidos == 4  # 3 cópias + 1 mesclado com o nome antigo
        assert not (saida / "Documentos Mesclados" / "Ana Souza.pdf").exists()
        assert (saida / "Documentos Mesclados" / "Ana Souza Santos.pdf").exists()
        assert (saida / "Documentos Mesclados" / "anotacoes pessoais.txt").exists()

    def test_arquivo_ausente_no_armazenamento_vira_pendencia(self, tmp_path):
        ana = ficha(nf(), planilha(), prestacao())
        servico = self._servico(tmp_path, [ana])
        (tmp_path / "storage" / f"{ana.documentos[0].hash_sha256}.pdf").unlink()

        resultado = servico.organizar()

        assert (resultado.documentos_copiados, resultado.pdfs_mesclados) == (2, 0)
        assert [p.motivo for p in resultado.pendencias] == [
            "Arquivo original não encontrado no armazenamento (storage).",
            "PDF mesclado não gerado: nem todos os documentos puderam ser copiados.",
        ]


class TestLibpq:
    def test_pasta_configurada_precisa_ter_a_libpq(self, tmp_path):
        assert localizar_pasta_da_libpq(tmp_path) is None
        (tmp_path / "libpq.dll").write_bytes(b"")
        assert localizar_pasta_da_libpq(tmp_path) == tmp_path

    def test_escolhe_a_maior_versao_instalada(self, tmp_path, monkeypatch):
        for versao in ("9.6", "18", "16"):
            pasta = tmp_path / versao / "bin"
            pasta.mkdir(parents=True)
            (pasta / "libpq.dll").write_bytes(b"")
        monkeypatch.setattr("app.core.libpq.PASTA_PADRAO_POSTGRESQL", Path(tmp_path))
        assert localizar_pasta_da_libpq() == tmp_path / "18" / "bin"
