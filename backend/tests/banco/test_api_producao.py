"""Funções de produção pela API, com PostgreSQL: salvar automaticamente, revisar antes de
salvar, painel, planilha .xlsx e organização em pastas."""

from io import BytesIO

import pymupdf
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.dependencies import obter_armazenamento, obter_fabrica_de_sessoes, obter_pipeline, obter_settings
from app.core.armazenamento import ArmazenamentoDeArquivos
from app.core.config import Settings
from app.extraction import criar_pipeline
from app.main import criar_app
from tests.api.conftest import OCR, TABELAS, TEXTOS
from tests.conftest import LeitorFalso, LeitorTabelasFalso, criar_pdf_com_texto

pytestmark = pytest.mark.banco

# PDFs de verdade (a organização mescla páginas), com os textos que os leitores falsos devolvem.
NF_ANA = criar_pdf_com_texto(["NF ANA"])
PLANILHA_ANA = criar_pdf_com_texto(["PLANILHA ANA"])
PRESTACAO_ANA = criar_pdf_com_texto(["PRESTACAO ANA"])
NF_ILEGIVEL = criar_pdf_com_texto(["NF ILEGIVEL"])
ARQUIVOS_ANA = [
    ("NF Ana Souza.pdf", NF_ANA),
    ("Planilha Comissão Ana Souza (VT Helena).pdf", PLANILHA_ANA),
    ("Prestação de Contas Ana Souza.pdf", PRESTACAO_ANA),
]


@pytest.fixture
def api(sessoes, tmp_path):
    textos = {PRESTACAO_ANA: next(iter(TEXTOS.values()))}
    ocr = {NF_ANA: next(iter(OCR.values()))}
    tabelas = {PLANILHA_ANA: next(iter(TABELAS.values()))}
    app = criar_app()
    pipeline = criar_pipeline(
        leitor_texto=LeitorFalso(textos),
        leitor_ocr=LeitorFalso(ocr),
        leitor_tabelas=LeitorTabelasFalso(tabelas),
        arquivo_correcoes=None,
    )
    app.dependency_overrides[obter_fabrica_de_sessoes] = lambda: sessoes
    app.dependency_overrides[obter_pipeline] = lambda: pipeline
    app.dependency_overrides[obter_armazenamento] = lambda: ArmazenamentoDeArquivos(tmp_path / "storage")
    app.dependency_overrides[obter_settings] = lambda: Settings(_env_file=None, pasta_documentos=tmp_path / "saida")
    with TestClient(app) as cliente:
        yield cliente


def _enviar(api, *arquivos, automatico=False):
    return api.post(
        "/api/v1/uploads",
        files=[("arquivos", (nome, conteudo, "application/pdf")) for nome, conteudo in arquivos],
        data={"importar_automaticamente": str(automatico).lower()},
    )


def test_lote_salvo_automaticamente_ao_terminar_a_leitura(api):
    lote = _enviar(api, *ARQUIVOS_ANA, automatico=True).json()

    assert lote["importar_automaticamente"] is True
    [importacao] = api.get("/api/v1/importacoes", params={"processamento_id": lote["id"]}).json()
    assert (importacao["clientes_criados"], importacao["documentos_criados"]) == (1, 3)
    painel = api.get("/api/v1/dashboard").json()
    assert (painel["total_clientes"], painel["clientes_completos"]) == (1, 1)


def test_sem_salvamento_automatico_nada_vai_para_o_banco(api):
    lote = _enviar(api, *ARQUIVOS_ANA).json()
    assert api.get("/api/v1/importacoes", params={"processamento_id": lote["id"]}).json() == []
    assert api.get("/api/v1/clientes").json() == []


def test_arquivo_revisado_no_lote_entra_com_as_correcoes_protegidas(api):
    _enviar(api, *ARQUIVOS_ANA, ("NF 0042.pdf", NF_ILEGIVEL))
    api.put(
        "/api/v1/processamentos/1/itens/4",
        json={"tipo": "NOTA_FISCAL", "nome_cliente": "ANA SOUZA", "data_emissao": "2024-05-02"},
    )

    api.post("/api/v1/importacoes", json={"processamento_id": 1})

    [ana] = api.get("/api/v1/clientes").json()
    revisada = next(d for d in ana["documentos"] if d["nome_arquivo"] == "NF 0042.pdf")
    assert revisada["campos_editados"] == ["cliente", "data_emissao"]
    assert revisada["campos_a_verificar"] == []
    historico = api.get(f"/api/v1/documentos/{revisada['id']}/historico").json()
    assert historico[0]["valor_novo"] == "Importado do processamento 1 (revisado à mão: data de emissão, nome do cliente)"


def test_arquivo_ja_salvo_nao_pode_ser_revisado_no_lote(api):
    _enviar(api, *ARQUIVOS_ANA, automatico=True)
    resposta = api.put("/api/v1/processamentos/1/itens/1", json={"tipo": "NOTA_FISCAL", "nome_cliente": "ANA SOUZA"})
    assert resposta.status_code == 409


def test_planilha_xlsx_para_download(api):
    _enviar(api, *ARQUIVOS_ANA, automatico=True)

    resposta = api.get("/api/v1/relatorios/clientes.xlsx")

    assert resposta.status_code == 200
    assert resposta.headers["content-disposition"].startswith('attachment; filename="relatorio_clientes_')
    linhas = list(load_workbook(BytesIO(resposta.content))["Clientes"].iter_rows(values_only=True))
    assert linhas[1][:5] == ("ANA SOUZA", None, "0001001-37.2014.8.05.0001", "HELENA", "Valor Total")


def test_organizar_documentos_em_pastas(api, tmp_path):
    _enviar(api, *ARQUIVOS_ANA, automatico=True)

    assert api.get("/api/v1/organizacoes/configuracao").json()["pasta"] == str(tmp_path / "saida")
    resultado = api.post("/api/v1/organizacoes", json={}).json()

    assert (resultado["documentos_copiados"], resultado["pdfs_mesclados"], resultado["pendencias"]) == (3, 1, [])
    mesclado = tmp_path / "saida" / "Documentos Mesclados" / "Ana Souza.pdf"
    with pymupdf.open(mesclado) as pdf:
        # Ordem do notebook: Prestação, Planilha, NF.
        assert [pagina.get_text().strip() for pagina in pdf] == ["PRESTACAO ANA", "PLANILHA ANA", "NF ANA"]
    assert (tmp_path / "saida" / "Documentos Verificados" / "Planilha Corretores" / "Planilha Comissão Ana Souza (VT Helena).pdf").exists()
