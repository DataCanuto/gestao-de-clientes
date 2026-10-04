"""Verifica a nova camada de extração contra os resultados do notebook e dos relatórios.

Uso (a partir de APP/backend):

    python -m scripts.verificar_paridade \
        --pasta "<pasta com os PDFs>" \
        --csv "../../RELATÓRIOS/tabela_final_clientes.csv" \
        --tabela-clientes "<planilha de controle>.xlsx" \
        --tesseract-cmd "C:/Program Files/Tesseract-OCR/tesseract.exe"

    --sem-ocr   pula o OCR das NFs (roda em segundos; não compara data_nf nem nomes de NF)
    --cache-ocr ARQUIVO.json  guarda o texto do OCR: a 1ª execução leva ~3 min, as
                seguintes, segundos (o arquivo contém dados de clientes: fora do repositório)

Seções do relatório:
1. resultado por arquivo (nome encontrado no conteúdo, situação, avisos);
2. paridade campo a campo com a tabela_final_clientes.csv do notebook, usando
   a CHAVE LEGADA (nome do arquivo), para isolar as regras de extração;
3. situação x coluna "Situação" do relatório Tabela Clientes (pelo processo);
4. agrupamento pela CHAVE NOVA (nome lido do conteúdo) e possíveis duplicidades;
5. números do dashboard.

O relatório sai só no terminal: contém nomes de clientes e não deve ser salvo
no repositório.
"""

import argparse
import csv
import json
import logging
import sys
import time
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from app.domain.consolidacao import ConsolidadorFichas, chave_legada_do_notebook
from app.domain.entities import FichaCliente, ResultadoExtracao
from app.domain.enums import Situacao, StatusExtracao, TipoDocumento
from app.domain.relatorio import gerar_resumo
from app.extraction import criar_pipeline
from app.extraction.readers import LeitorComCache, PdfOcrReader

# Textos que o notebook gravava no lugar de "sem valor".
MARCADORES_DE_AUSENCIA = {
    "",
    "nan",
    "Processo não encontrado",
    "Número não encontrado",
    "Situação não encontrada",
    "Corretor não encontrado",
    "Verificar prestação",
    "Nome não encontrado",
    "Erro no processamento",
}
CAMPOS_COMPARADOS = ("data_nf", "n_processo", "corretor", "nome_encontrado")
SITUACAO_DO_RELATORIO = {
    "Valor Total": Situacao.VALOR_TOTAL,
    "Valor Final": Situacao.VALOR_FINAL,
    "Incontroverso": Situacao.VALOR_INCONTROVERSO,
    "Acordo": Situacao.ACORDO,
}
MAXIMO_EXEMPLOS = 15


class LeitorVazio:
    """Substitui o OCR quando --sem-ocr é usado."""

    def ler_texto(self, conteudo: bytes) -> str:
        return ""


class CacheEmArquivoJson(dict):
    """Dicionário que se grava num arquivo JSON a cada item novo.

    Usado com o LeitorComCache para que a 2ª execução do script não repita o OCR.
    O arquivo contém texto dos documentos (nomes de clientes): mantenha-o FORA
    do repositório.
    """

    def __init__(self, caminho: Path) -> None:
        super().__init__(json.loads(caminho.read_text(encoding="utf-8")) if caminho.exists() else {})
        self._caminho = caminho

    def __setitem__(self, chave: str, valor: str) -> None:
        super().__setitem__(chave, valor)
        self._caminho.write_text(json.dumps(self, ensure_ascii=False), encoding="utf-8")


@dataclass
class Divergencia:
    cliente: str
    campo: str
    esperado: object
    novo: object


def main() -> int:
    argumentos = _ler_argumentos()
    logging.basicConfig(level=logging.WARNING, format="%(levelname)s %(name)s: %(message)s")

    if argumentos.sem_ocr:
        leitor_ocr = LeitorVazio()
    elif argumentos.cache_ocr:
        leitor_ocr = LeitorComCache(PdfOcrReader(argumentos.tesseract_cmd), CacheEmArquivoJson(argumentos.cache_ocr))
    else:
        leitor_ocr = None  # OCR real, sem cache
    pipeline = criar_pipeline(tesseract_cmd=argumentos.tesseract_cmd, leitor_ocr=leitor_ocr)
    resultados = _processar_pasta(pipeline, argumentos.pasta)
    _relatorio_por_arquivo(resultados)

    legado = ConsolidadorFichas(chave_legada_do_notebook, chave_reserva=None, reconciliar_ocr=False).consolidar(
        resultados
    )
    _relatorio_paridade_notebook(legado.fichas, _ler_csv(argumentos.csv), comparar_data=not argumentos.sem_ocr)

    novo = ConsolidadorFichas().consolidar(resultados)
    if argumentos.tabela_clientes:
        _relatorio_situacao(novo.fichas, argumentos.tabela_clientes)
    _relatorio_agrupamento(legado, novo)
    _relatorio_dashboard(novo.fichas)
    return 0


def _ler_argumentos() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--pasta", type=Path, required=True, help="pasta com os PDFs originais")
    parser.add_argument("--csv", type=Path, required=True, help="tabela_final_clientes.csv do notebook")
    parser.add_argument("--tabela-clientes", type=Path, help="Planilha de controle do escritório (situação real)")
    parser.add_argument("--tesseract-cmd", default=None, help="caminho do tesseract.exe")
    parser.add_argument("--sem-ocr", action="store_true", help="não roda OCR nas NFs")
    parser.add_argument("--cache-ocr", type=Path, help="arquivo JSON de cache do OCR (fora do repositório!)")
    return parser.parse_args()


def _processar_pasta(pipeline, pasta: Path) -> list[ResultadoExtracao]:
    arquivos = sorted((p for p in pasta.iterdir() if p.is_file()), key=lambda p: p.name.upper())
    resultados = []
    inicio = time.perf_counter()
    for posicao, caminho in enumerate(arquivos, start=1):
        resultados.append(pipeline.processar(caminho.name, caminho.read_bytes()))
        if posicao % 100 == 0:
            print(f"  ... {posicao}/{len(arquivos)} arquivos ({time.perf_counter() - inicio:.0f}s)", flush=True)
    print(f"Processados {len(arquivos)} arquivos em {time.perf_counter() - inicio:.1f}s")
    return resultados


# --------------------------------------------------------------------------- 1


def _relatorio_por_arquivo(resultados: list[ResultadoExtracao]) -> None:
    _titulo("1. RESULTADO POR ARQUIVO")
    print("Status:", dict(Counter(r.status.value for r in resultados)))
    for resultado in resultados:
        if resultado.status is not StatusExtracao.SUCESSO:
            print(f"  {resultado.status.value} {resultado.nome_arquivo}: {resultado.erro or resultado.avisos}")

    sucesso = [r for r in resultados if r.status is StatusExtracao.SUCESSO]
    print("\nPor tipo: nome encontrado NO CONTEÚDO | situação")
    for tipo, total in sorted(Counter(r.tipo for r in sucesso).items()):
        do_tipo = [r for r in sucesso if r.tipo is tipo]
        com_nome = sum(1 for r in do_tipo if r.nome_pelo_conteudo)
        situacoes = Counter(r.dados.situacao.value if r.dados.situacao else "-" for r in do_tipo)
        print(f"  {tipo.value:18} nome {com_nome:3}/{total:3} | {dict(situacoes)}")

    sem_nome = [r.nome_arquivo for r in sucesso if not r.nome_pelo_conteudo]
    print(f"\nSem nome no conteúdo ({len(sem_nome)}):")
    for nome in sem_nome[:MAXIMO_EXEMPLOS]:
        print(f"  {nome}")

    avisos = Counter(aviso.split(":")[0].split("'")[0] for r in sucesso for aviso in r.avisos)
    print("\nAvisos (agrupados):")
    for aviso, quantidade in avisos.most_common():
        print(f"  {quantidade:3} x {aviso}")


# --------------------------------------------------------------------------- 2


def _ler_csv(caminho: Path) -> dict[str, dict[str, str | None]]:
    with caminho.open(encoding="utf-8-sig", newline="") as arquivo:
        return {
            linha["nome"]: {campo: _sem_marcador(linha[campo]) for campo in CAMPOS_COMPARADOS}
            for linha in csv.DictReader(arquivo)
        }


def _sem_marcador(valor: str | None) -> str | None:
    return None if valor is None or valor.strip() in MARCADORES_DE_AUSENCIA else valor.strip()


def _valores_como_no_notebook(ficha: FichaCliente) -> dict[str, str | None]:
    """O notebook guardava só o PRIMEIRO documento de cada tipo. A ficha nova guarda
    todos; para comparar com o CSV antigo, reproduzimos aqui a regra antiga."""
    [nf] = ficha.documentos_do_tipo(TipoDocumento.NOTA_FISCAL)[:1] or [None]
    [planilha] = ficha.documentos_do_tipo(TipoDocumento.PLANILHA_COMISSAO)[:1] or [None]
    [prestacao] = ficha.documentos_do_tipo(TipoDocumento.PRESTACAO_CONTAS)[:1] or [None]
    return {
        "data_nf": _formatar_data(nf.dados.data_emissao) if nf else None,
        "n_processo": prestacao.dados.numero_processo if prestacao else None,
        "corretor": planilha.dados.corretor if planilha else None,
        "nome_encontrado": prestacao.dados.nome_encontrado if prestacao else None,
    }


def _formatar_data(valor: date | None) -> str | None:
    return valor.strftime("%d/%m/%Y") if valor else None


def _relatorio_paridade_notebook(fichas: list[FichaCliente], csv_notebook: dict, comparar_data: bool) -> None:
    _titulo("2. PARIDADE COM O NOTEBOOK (chave legada = nome do arquivo)")
    novas = {ficha.chave: ficha for ficha in fichas}
    so_no_csv = sorted(csv_notebook.keys() - novas.keys())
    so_no_novo = sorted(novas.keys() - csv_notebook.keys())
    print(f"Clientes: {len(novas)} (novo) x {len(csv_notebook)} (notebook)")
    print(f"Só no CSV do notebook ({len(so_no_csv)}): {so_no_csv}")
    print(f"Só na extração nova ({len(so_no_novo)}): {so_no_novo}")

    divergencias: list[Divergencia] = []
    for nome in sorted(novas.keys() & csv_notebook.keys()):
        valores_novos = _valores_como_no_notebook(novas[nome])
        for campo in CAMPOS_COMPARADOS:
            if campo == "data_nf" and not comparar_data:
                continue
            if valores_novos[campo] != csv_notebook[nome][campo]:
                divergencias.append(Divergencia(nome, campo, csv_notebook[nome][campo], valores_novos[campo]))

    print(f"\nDivergências de campo: {len(divergencias)} {dict(Counter(d.campo for d in divergencias))}")
    for d in divergencias:
        print(f"  [{d.campo}] {d.cliente}: notebook={d.esperado!r} novo={d.novo!r}")


# --------------------------------------------------------------------------- 3


def _relatorio_situacao(fichas: list[FichaCliente], caminho_tabela: Path) -> None:
    _titulo("3. SITUAÇÃO x RELATÓRIO 'Tabela Clientes' (cruzamento pelo nº do processo)")
    from openpyxl import load_workbook  # dependência só de desenvolvimento

    planilha = load_workbook(caminho_tabela, read_only=True)["Table1"]
    linhas = planilha.iter_rows(values_only=True)
    cabecalho = next(linhas)
    indice = {coluna: posicao for posicao, coluna in enumerate(cabecalho)}
    situacao_real = {
        linha[indice["Número do Processo"]]: linha[indice["Situação"]]
        for linha in linhas
        if linha[indice["Número do Processo"]]
    }

    # Um cliente pode ter vários processos e várias planilhas: vale se ALGUMA
    # planilha do cliente tem a situação que o relatório registra para o processo.
    cruzamentos = [
        (ficha, situacao_real[processo])
        for ficha in fichas
        for processo in ficha.numeros_processo
        if processo in situacao_real
    ]
    comparaveis = [(f, real) for f, real in cruzamentos if real in SITUACAO_DO_RELATORIO]
    iguais = [(f, real) for f, real in comparaveis if SITUACAO_DO_RELATORIO[real] in f.situacoes]
    print(f"Processos cruzados: {len(cruzamentos)}; com situação do enum no relatório: {len(comparaveis)}")
    print(f"Situação do relatório presente nas planilhas do cliente: {len(iguais)}/{len(comparaveis)}")
    for ficha, real in comparaveis:
        if (ficha, real) not in iguais:
            extraidas = [s.value for s in ficha.situacoes] or "sem planilha"
            print(f"  {ficha.nome}: planilhas={extraidas} relatório={real!r}")
    fora = Counter(real for _, real in cruzamentos if real not in SITUACAO_DO_RELATORIO)
    print(f"Situações do relatório fora do enum: {dict(fora)}")


# --------------------------------------------------------------------------- 4


def _relatorio_agrupamento(legado, novo) -> None:
    _titulo("4. AGRUPAMENTO PELA CHAVE NOVA (nome lido do conteúdo)")
    print(f"Clientes: {len(novo.fichas)} (chave do conteúdo) x {len(legado.fichas)} (chave legada)")
    pelo_arquivo = [f.nome for f in novo.fichas if f.origem_chave.value == "NOME_ARQUIVO"]
    print(f"Fichas identificadas só pelo nome do arquivo (reserva): {len(pelo_arquivo)} {pelo_arquivo[:MAXIMO_EXEMPLOS]}")
    print(f"Documentos sem cliente: {[d.nome_arquivo for d in novo.documentos_sem_cliente]}")
    print(f"\nNomes de OCR reconciliados com a camada de texto: {len(novo.reconciliacoes_ocr)}")
    for reconciliacao in sorted(novo.reconciliacoes_ocr, key=lambda r: r.similaridade)[:MAXIMO_EXEMPLOS]:
        print(f"  {reconciliacao.similaridade:.2f}  {reconciliacao.chave_lida!r} -> {reconciliacao.chave_adotada!r}")
    print(f"\nPossíveis duplicidades para revisão: {len(novo.possiveis_duplicidades)}")
    for duplicidade in novo.possiveis_duplicidades:
        print(f"  {duplicidade.chave_a}  <->  {duplicidade.chave_b}\n      {duplicidade.motivo}")


# --------------------------------------------------------------------------- 5


def _relatorio_dashboard(fichas: list[FichaCliente]) -> None:
    _titulo("5. NÚMEROS DO DASHBOARD")
    resumo = gerar_resumo(fichas)
    print(f"Total de clientes: {resumo.total_clientes}")
    print(f"Total de documentos: { {t.value: n for t, n in resumo.total_documentos.items()} }")
    print(f"Clientes sem o documento: { {t.value: n for t, n in resumo.clientes_sem_documento.items()} }")
    print(f"Clientes com >1 documento: { {t.value: n for t, n in resumo.clientes_com_varios_documentos.items()} }")
    print(f"Clientes com documentos faltantes: {resumo.clientes_com_documentos_faltantes}")
    print(f"Documentos a verificar: {resumo.documentos_a_verificar}")
    print(f"Planilhas por situação: { {(s.value if s else None): n for s, n in resumo.planilhas_por_situacao.items()} }")


def _titulo(texto: str) -> None:
    print(f"\n{'=' * 78}\n{texto}\n{'=' * 78}")


if __name__ == "__main__":
    sys.exit(main())
